from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation

from baseline import load_production_baseline

from astromol.census import CensusView
from astromol.database import Database
from astromol.slides import (
    LEGACY_SLIDE_HEIGHT_IN,
    LEGACY_SLIDE_WIDTH_IN,
    PPD_TITLE,
    build_ppd_detection_slide_layout,
    build_molecule_slide_layout,
    molecule_slide_report,
    powerpoint_formula_runs,
    selected_molecule_slide_entries,
    slide_version_label,
    write_molecule_slide,
    write_ppd_detection_slide,
    write_molecule_slide_report,
)


def run_text(runs):
    """Return combined run text for a PowerPoint formula tokenization."""
    return "".join(run.text for run in runs)


def baselines(runs):
    """Return non-normal baseline assignments by run text."""
    return {
        run.text: run.baseline
        for run in runs
        if run.baseline != "normal"
    }


db = Database()
counts_2026 = load_production_baseline()["regression_counts"][
    "slides_molecule_slide_2026"
]
ppd_counts_2026 = load_production_baseline()["regression_counts"][
    "slides_ppd_detection_slide_2026"
]
view_2021 = CensusView.for_census(db, "2021")
view_2026 = CensusView.for_census(db, "2026")

entries_2021 = selected_molecule_slide_entries(view_2021)
assert len(entries_2021) == 240
assert len({entry.label for entry in entries_2021}) == 240
assert {entry.label for entry in entries_2021} == {
    molecule.label
    for molecule in view_2021.ism_molecules()
    if molecule.isotopologue_of is None
}
assert all(entry.molecule.isotopologue_of is None for entry in entries_2021)

layout_2021 = build_molecule_slide_layout(view_2021)
assert layout_2021.total == 240
assert [len(group.molecules) for group in layout_2021.groups] == [
    41,
    45,
    31,
    31,
    23,
    15,
    15,
    14,
    6,
    6,
    5,
    8,
]
assert layout_2021.warnings == ()
assert len(layout_2021.molecules) == layout_2021.total

for group in layout_2021.groups:
    assert group.spec.label_box.right <= LEGACY_SLIDE_WIDTH_IN
    assert group.spec.label_box.bottom <= LEGACY_SLIDE_HEIGHT_IN
    for column in group.columns:
        assert column.box.right <= LEGACY_SLIDE_WIDTH_IN
        assert column.box.bottom <= LEGACY_SLIDE_HEIGHT_IN

report_2021 = molecule_slide_report(layout_2021)
assert "- total molecules: 240" in report_2021
assert "- 13+ Atoms: 8 molecule(s); columns = 4, 4" in report_2021
assert "- none" in report_2021

layout_2026 = build_molecule_slide_layout(view_2026)
assert layout_2026.total == counts_2026["total"]
assert len(layout_2026.warnings) == counts_2026["legacy_warning_count"]
assert any("13+ Atoms" in warning for warning in layout_2026.warnings)

balanced_2026 = build_molecule_slide_layout(view_2026, profile="balanced")
assert balanced_2026.total == counts_2026["total"]
assert balanced_2026.profile == "balanced"
assert balanced_2026.molecule_font_pt == counts_2026["balanced_molecule_font_pt"]
assert balanced_2026.warnings == ()
assert [len(group.molecules) for group in balanced_2026.groups] == counts_2026[
    "balanced_group_molecule_counts"
]
assert [len(group.columns) for group in balanced_2026.groups] == counts_2026[
    "balanced_group_column_counts"
]
assert len(balanced_2026.molecules) == balanced_2026.total

for group in balanced_2026.groups:
    assert group.spec.label_box.right <= LEGACY_SLIDE_WIDTH_IN
    assert group.spec.label_box.bottom <= LEGACY_SLIDE_HEIGHT_IN
    for column in group.columns:
        assert column.box.right <= LEGACY_SLIDE_WIDTH_IN
        assert column.box.bottom <= LEGACY_SLIDE_HEIGHT_IN
        assert column.overflow == 0

c3hcch_runs = powerpoint_formula_runs("c-C3HCCH")
assert run_text(c3hcch_runs) == "c-C3HCCH"
assert c3hcch_runs[0].italic is True
assert baselines(c3hcch_runs)["3"] == "subscript"

anion_runs = powerpoint_formula_runs("C6H-")
assert run_text(anion_runs) == "C6H-"
assert baselines(anion_runs)["6"] == "subscript"
assert baselines(anion_runs)["-"] == "superscript"

positional_runs = powerpoint_formula_runs("1-C5H5CN")
assert run_text(positional_runs) == "1-C5H5CN"
assert positional_runs[0].baseline == "normal"
assert positional_runs[1].baseline == "normal"

isotope_runs = powerpoint_formula_runs("[13C]CC")
assert run_text(isotope_runs) == "13CCC"
assert isotope_runs[0].baseline == "superscript"
assert isotope_runs[1].baseline == "normal"

assert slide_version_label("development") == "development"
assert slide_version_label("development (git abc123)") == "development (git abc123)"
assert slide_version_label("2026.0.0") == "v2026.0.0"

ppd_2026 = build_ppd_detection_slide_layout(view_2026)
assert ppd_2026.title == PPD_TITLE
assert ppd_2026.total == ppd_counts_2026["total"]
assert ppd_2026.profile == "compact"
assert ppd_2026.detection_type == "ppd"
assert ppd_2026.include_isotopologues is True
assert ppd_2026.molecule_font_pt == ppd_counts_2026["molecule_font_pt"]
assert ppd_2026.warnings == ()
assert [group.spec.label for group in ppd_2026.groups] == ppd_counts_2026["group_labels"]
assert [len(group.molecules) for group in ppd_2026.groups] == ppd_counts_2026[
    "group_molecule_counts"
]
assert any(entry.molecule.isotopologue_of for entry in ppd_2026.molecules)

with TemporaryDirectory() as tmp:
    output_dir = Path(tmp)
    pptx_path = write_molecule_slide(
        view_2021,
        output_dir / "astro_molecules_2021.pptx",
        last_updated="2026-05-11",
    )
    report_path = write_molecule_slide_report(
        layout_2021,
        output_dir / "astro_molecules_2021_report.md",
    )

    assert pptx_path.exists()
    assert pptx_path.stat().st_size > 0
    assert report_path.exists()
    assert "- total molecules: 240" in report_path.read_text()

    presentation = Presentation(pptx_path)
    assert len(presentation.slides) == 1
    slide_text = "\n".join(
        shape.text
        for shape in presentation.slides[0].shapes
        if hasattr(shape, "text")
    )
    assert "Known Interstellar Molecules" in slide_text
    assert "240 Molecules" in slide_text
    assert "Last Updated: 11 May 2026" in slide_text
    assert "c-C3HCCH" in slide_text

    balanced_path = write_molecule_slide(
        view_2026,
        output_dir / "astro_molecules_2026_balanced.pptx",
        profile="balanced",
        last_updated="2026-05-11",
    )
    balanced_report_path = write_molecule_slide_report(
        balanced_2026,
        output_dir / "astro_molecules_2026_balanced_report.md",
    )
    assert balanced_path.exists()
    assert balanced_path.stat().st_size > 0
    assert "- profile: `balanced`" in balanced_report_path.read_text()

    balanced_presentation = Presentation(balanced_path)
    assert len(balanced_presentation.slides) == 1
    balanced_text = "\n".join(
        shape.text
        for shape in balanced_presentation.slides[0].shapes
        if hasattr(shape, "text")
    )
    assert "Known Interstellar Molecules" in balanced_text
    assert counts_2026["count_text"] in balanced_text
    assert "Last Updated: 11 May 2026" in balanced_text

    ppd_path = write_ppd_detection_slide(
        view_2026,
        output_dir / "ppd_molecules_2026.pptx",
        last_updated="2026-05-11",
    )
    ppd_presentation = Presentation(ppd_path)
    assert len(ppd_presentation.slides) == 1
    ppd_text = "\n".join(
        shape.text
        for shape in ppd_presentation.slides[0].shapes
        if hasattr(shape, "text")
    )
    assert PPD_TITLE in ppd_text
    assert ppd_counts_2026["count_text"] in ppd_text
    assert "Last Updated: 11 May 2026" in ppd_text
    assert "13CO" in ppd_text

print("Molecule slide generation verification passed")
